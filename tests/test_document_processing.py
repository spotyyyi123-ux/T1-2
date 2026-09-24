import asyncio
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

import docx
import httpx

from app.services import document_processing as pipeline
from app.services.extraction_format import AIServiceError, display_attributes, normalize_ai_attributes
from app.core.json_codec import dumps
from app.services.text_extraction import extract_text


class ChunkTests(unittest.TestCase):
    def test_every_character_is_covered_without_gaps(self):
        for text in ('', 'а', 'a' * 3001, ('Условие договора.\n\n' * 350) + 'ХВОСТ'):
            for size, overlap in ((100, 0), (100, 15), (150, 40)):
                with self.subTest(length=len(text), size=size, overlap=overlap):
                    chunks = pipeline.split_text(text, size, overlap, 1000)
                    rebuilt = ''
                    cursor = 0
                    for chunk in chunks:
                        self.assertEqual(chunk.text, text[chunk.start:chunk.end])
                        self.assertLessEqual(len(chunk.text), size)
                        self.assertLessEqual(chunk.start, cursor)
                        self.assertGreater(chunk.end, cursor)
                        rebuilt += chunk.text[cursor - chunk.start:]
                        cursor = chunk.end
                    self.assertEqual(rebuilt, text)

    def test_limits_and_invalid_configuration_fail_explicitly(self):
        with self.assertRaisesRegex(AIServiceError, 'лимит 2 частей'):
            pipeline.split_text('a' * 500, 100, 10, 2)
        for settings in ((0, 0, 10), (100, 50, 10), (100, -1, 10), (100, 10, 0)):
            with self.assertRaises(AIServiceError):
                pipeline.split_text('text', *settings)

    def test_location_of_chunk_start_and_next_page(self):
        text = '[Страница 1]\n' + 'а' * 100 + '\n[Страница 2]\n' + 'б' * 100
        chunk = pipeline.TextChunk(40, len(text), text[40:])
        self.assertEqual(pipeline.chunk_locations(text, chunk), ['Страница 1', 'Страница 2'])

    def test_merge_normalizes_and_deduplicates_without_summing(self):
        result = pipeline.merge_results([
            {'amount': '1 250,50 RUB', 'contract_date': '20.09.2026', 'customer': 'ООО А'},
            {'amount': Decimal('1250.50'), 'contract_date': '2026-09-20', 'customer': 'ооо а'},
            {'contractor': 'ООО Б'},
        ])
        self.assertEqual(result['amount'], Decimal('1250.50'))
        self.assertEqual(result['contract_date'], '2026-09-20')
        self.assertEqual(result['customer'], 'ООО А')
        self.assertEqual(result['contractor'], 'ООО Б')
        self.assertEqual(result['unparsed_values'], {})

    def test_conflicts_survive_normalization_and_are_visible(self):
        result = pipeline.merge_results([
            {'amount': 100, 'contract_date': '2026-09-20', 'payment_terms': 'Аванс 20%'},
            {'amount': 200, 'contract_date': '2026-09-21', 'payment_terms': 'Аванс 50%'},
        ])
        for field in ('amount', 'contract_date', 'payment_terms'):
            self.assertIsNone(result[field])
            self.assertIn('части 1:', result['unparsed_values'][field])
            self.assertIn('части 2:', result['unparsed_values'][field])
            self.assertIn('нужна проверка', display_attributes(result)[field])
        self.assertEqual(normalize_ai_attributes(result), result)

    def test_rejected_value_cannot_be_overruled_by_another_chunk(self):
        result = pipeline.merge_results([{'amount': 100}, {'amount': '100 или 200'}])
        self.assertIsNone(result['amount'])
        self.assertIn('100 или 200', result['unparsed_values']['amount'])

    def test_currency_conflict_also_blocks_amount(self):
        result = pipeline.merge_results([{'amount': 100, 'currency': 'RUB'}, {'amount': 100, 'currency': 'USD'}])
        self.assertIsNone(result['amount'])
        self.assertIsNone(result['currency'])


class PipelineTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        for key, value in {'POLZA_API_KEY': 'test-only', 'POLZA_API_URL': 'https://provider.invalid',
                           'MAX_PROMPT_CHARS': 100, 'CHUNK_OVERLAP_CHARS': 15,
                           'MAX_DOCUMENT_CHUNKS': 40}.items():
            p = patch.object(pipeline, key, value)
            p.start()
            self.addCleanup(p.stop)
        self.client = AsyncMock()
        p = patch.object(pipeline.httpx, 'AsyncClient')
        self.factory = p.start()
        self.factory.return_value.__aenter__.return_value = self.client
        self.addCleanup(p.stop)

    @staticmethod
    def reply(data, status=200, finish='stop'):
        return httpx.Response(status, request=httpx.Request('POST', 'https://provider.invalid'),
                              json={'choices': [{'finish_reason': finish, 'message': {'content': dumps(data)}}]})

    async def test_tail_beyond_old_limit_is_sent_and_extracted(self):
        text = '[Абзац 1]\n' + ('Длинное условие.\n' * 1000) + '\nTOTAL_AT_END'
        progress = []

        async def respond(*args, **kwargs):
            content = kwargs['json']['messages'][-1]['content']
            return self.reply({'amount': 12345 if 'TOTAL_AT_END' in content else None})

        self.client.post.side_effect = respond
        with patch.object(pipeline, 'MAX_PROMPT_CHARS', 1000):
            result = await pipeline.fetch_ai_attributes(text, lambda done, total: progress.append((done, total)))
        self.assertEqual(result['amount'], 12345)
        self.assertEqual(result['processing']['text_chars'], len(text))
        total = result['processing']['chunks_total']
        self.assertEqual(progress, [(n, total) for n in range(total + 1)])
        self.assertEqual(self.client.post.await_count, total)
        self.assertEqual(result['processing']['chunks'][-1]['end'], len(text))
        self.assertEqual(normalize_ai_attributes(result), result)

    async def test_failure_in_middle_does_not_return_partial_result(self):
        self.client.post.side_effect = [self.reply({'amount': 100}), self.reply({}, status=503), self.reply({}, status=503)]
        progress = []
        with patch.object(pipeline.asyncio, 'sleep', new=AsyncMock()):
            with self.assertRaisesRegex(AIServiceError, 'части 2/'):
                await pipeline.fetch_ai_attributes('x' * 350, lambda done, total: progress.append(done))
        self.assertEqual(progress, [0, 1])
        self.assertEqual(self.client.post.await_count, 3)

    async def test_one_retry_for_transient_error(self):
        self.client.post.side_effect = [self.reply({}, status=429), self.reply({'amount': 100})]
        with patch.object(pipeline.asyncio, 'sleep', new=AsyncMock()):
            result = await pipeline.fetch_ai_attributes('short')
        self.assertEqual(result['amount'], 100)
        self.assertEqual(self.client.post.await_count, 2)

    async def test_local_ambiguity_is_not_lost_when_other_chunk_has_a_value(self):
        self.client.post.side_effect = [self.reply({'amount': 100}), self.reply({
            'amount': None, 'uncertain_fields': {'amount': 'Общие суммы 100 и 200'}
        })]
        result = await pipeline.fetch_ai_attributes('x' * 150)
        self.assertIsNone(result['amount'])
        self.assertIn('Общие суммы 100 и 200', result['unparsed_values']['amount'])

    async def test_invalid_or_truncated_response_is_not_success(self):
        for response in (self.reply({'amount': 10}, finish='length'), self.reply({}),
                         self.reply([]), self.reply({'amount': 10}, status=401)):
            with self.subTest(response=response):
                self.client.post.reset_mock()
                self.client.post.return_value = response
                with self.assertRaises(AIServiceError):
                    await pipeline.fetch_ai_attributes('short')
                self.assertEqual(self.client.post.await_count, 1)

    async def test_oversized_and_empty_documents_make_no_requests(self):
        for text in ('   ', 'x' * 5000):
            with self.assertRaises(AIServiceError):
                await pipeline.fetch_ai_attributes(text)
        self.client.post.assert_not_called()

    async def test_progress_callback_can_cancel_subsequent_requests(self):
        self.client.post.return_value = self.reply({'amount': 100})
        def progress(done, total):
            if done == 1:
                raise RuntimeError('document deleted')
        with self.assertRaisesRegex(RuntimeError, 'document deleted'):
            await pipeline.fetch_ai_attributes('x' * 350, progress)
        self.assertEqual(self.client.post.await_count, 1)


class TextExtractionTests(unittest.TestCase):
    def test_docx_preserves_table_order_and_nested_text(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'test.docx'
            document = docx.Document()
            document.add_paragraph('BEFORE_TABLE')
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = 'INSIDE_TABLE'
            nested = table.cell(0, 0).add_table(rows=1, cols=1)
            nested.cell(0, 0).text = 'NESTED_TABLE'
            document.add_paragraph('AFTER_TABLE')
            document.sections[0].header.paragraphs[0].text = 'HEADER_VALUE'
            document.save(path)
            text = extract_text(str(path))
        self.assertLess(text.index('BEFORE_TABLE'), text.index('INSIDE_TABLE'))
        self.assertLess(text.index('INSIDE_TABLE'), text.index('NESTED_TABLE'))
        self.assertLess(text.index('NESTED_TABLE'), text.index('AFTER_TABLE'))
        self.assertIn('HEADER_VALUE', text)
        self.assertIn('[Абзац 1]', text)

    def test_pdf_includes_every_text_page_and_keeps_page_numbers(self):
        pages = [MagicMock(), MagicMock(), MagicMock()]
        for page, value in zip(pages, ['FIRST', '', 'LAST']):
            page.extract_text.return_value = value
        with tempfile.NamedTemporaryFile(suffix='.pdf') as file:
            with patch('app.services.text_extraction.PdfReader', return_value=SimpleNamespace(pages=pages)):
                text = extract_text(file.name)
        self.assertEqual(text, '[Страница 1]\nFIRST\n\n[Страница 3]\nLAST')

    def test_empty_docx_stays_empty_not_just_location_markers(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'empty.docx'
            docx.Document().save(path)
            self.assertEqual(extract_text(str(path)), '')


if __name__ == '__main__':
    unittest.main()
